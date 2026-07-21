from pathlib import Path
from typing import List

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.services.cycle_service import CycleService


class ExcelParser:

    @staticmethod
    def _int(value):
        if value is None or value == "":
            return None
        try:
            return int(value)
        except Exception:
            return None

    @staticmethod
    def _float(value):
        if value is None or value == "":
            return None
        try:
            return float(value)
        except Exception:
            return None

    @staticmethod
    def _str(value):
        if value is None:
            return None
        return str(value).strip()

    @staticmethod
    def parse(
        db: Session,
        *,
        video_run_id: int,
        excel_path: str,
    ) -> List:

        excel_path = Path(excel_path)

        if not excel_path.exists():
            raise FileNotFoundError(excel_path)

        workbook = load_workbook(
            excel_path,
            data_only=True,
        )

        sheet = workbook.active

        headers = []

        for cell in sheet[1]:
            headers.append(str(cell.value).strip())

        saved_cycles = []

        for row in sheet.iter_rows(min_row=2, values_only=True):

            if not any(row):
                continue

            record = dict(zip(headers, row))

            cycle = CycleService.save_cycle(
                db=db,

                video_run_id=video_run_id,

                cycle_number=ExcelParser._int(
                    record.get("Cycle No.")
                ),

                start_frame=None,

                end_frame=None,

                duration_seconds=None,

                final_verdict=ExcelParser._str(
                    record.get("Final Verdict")
                ) or "UNKNOWN",

                output_video_path=ExcelParser._str(
                    record.get("Output Path")
                ) or "",

                tube_blue=ExcelParser._str(
                    record.get("tube_blue (Yellow)")
                ),

                transition_middle=ExcelParser._str(
                    record.get("trans_mid_tube (Blue)")
                ),

                transition_end=ExcelParser._str(
                    record.get("trans_end_tube (Pink)")
                ),

                detected_sequence=ExcelParser._str(
                    record.get("Detected Sequence")
                ),

                tube_order_result=ExcelParser._str(
                    record.get("Tube Order Result")
                ),

                anomaly_ratio=ExcelParser._float(
                    record.get("Anomaly Ratio")
                ),

                ok_votes=ExcelParser._int(
                    record.get("OK Votes")
                ),

                anomaly_votes=ExcelParser._int(
                    record.get("Anomaly Votes")
                ),

                total_frames=ExcelParser._int(
                    record.get("Total Frames")
                ),

                warmup_frames=ExcelParser._int(
                    record.get("Warmup Frames")
                ),

                inference_frames=ExcelParser._int(
                    record.get("Live Infer Frames")
                ),

                average_fps=ExcelParser._float(
                    record.get("Avg FPS")
                ),
            )

            saved_cycles.append(cycle)

        workbook.close()

        return saved_cycles